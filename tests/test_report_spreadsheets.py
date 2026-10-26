"""Verify exported bytes with real CSV/openpyxl readers, without Qt or devices."""
import csv
import importlib.util
import io
from pathlib import Path
import tempfile
import unittest

DESKTOP_EXPORTS_AVAILABLE = all(importlib.util.find_spec(name) for name in ('openpyxl', 'pandas'))
if DESKTOP_EXPORTS_AVAILABLE:
    import openpyxl
    import pandas as pd

    from desktop_ui.reports import export_frame, write_spreadsheet


@unittest.skipUnless(DESKTOP_EXPORTS_AVAILABLE, 'Desktop export dependencies not installed')
class ReportSpreadsheetTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)

    def path(self, suffix):
        return Path(self.directory.name) / ('report.' + suffix)

    def test_csv_neutralizes_text_without_changing_numeric_values_or_source(self):
        attacks = ['=1+1', '+SUM(1,2)', '-1+2', '@SUM(1)', '  =1+1', '\t=1+1',
                   '\r=1+1', '\n=1+1', '\ufeff=1+1', '\x00=1+1', '\tplain', '\ufeff \x00\t=1+1']
        frame = pd.DataFrame({'Worker': attacks, 'Delta': [-1.25] * len(attacks)})
        original = frame.copy(deep=True)
        write_spreadsheet(frame, self.path('csv'))
        raw = self.path('csv').read_bytes()
        self.assertTrue(raw.startswith(b'\xef\xbb\xbf'))
        rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'), newline='')))
        for row, attack in zip(rows[1:], attacks):
            self.assertEqual(row, ["'" + attack, '-1.25'])
        pd.testing.assert_frame_equal(frame, original)

    def test_xlsx_preserves_text_as_text_and_numbers_as_numbers(self):
        texts = ['=1+1', '+SUM(1,2)', '-1+2', '@SUM(1)', '#N/A', 'Клиент & пул', 'a,b"c']
        frame = pd.DataFrame({'=header': texts, 'Delta': [-1.25] * len(texts)})
        write_spreadsheet(frame, self.path('xlsx'))
        workbook = openpyxl.load_workbook(self.path('xlsx'), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active['A1'].data_type, 's')
        for row, text in zip(workbook.active.iter_rows(min_row=2), texts):
            self.assertEqual((row[0].value, row[0].data_type), (text, 's'))
            self.assertEqual((row[1].value, row[1].data_type), (-1.25, 'n'))

    def test_sort_and_column_selection_remain_separate_from_writer(self):
        rows = [{'IP': '192.0.2.10', 'Model': '=second'}, {'IP': '192.0.2.2', 'Model': 'first'}]
        frame = export_frame(rows, ['Model'], 'IP')
        self.assertEqual(frame.iloc[:, 0].tolist(), ['first', '=second'])
        write_spreadsheet(frame, self.path('csv'))
        self.assertEqual(list(csv.reader(io.StringIO(self.path('csv').read_text(encoding='utf-8-sig'))))[2][0], "'=second")

    def test_maximum_xlsx_text_is_preserved_and_oversize_is_explicitly_rejected(self):
        text = '=' + 'x' * 32766
        write_spreadsheet(pd.DataFrame({'Worker': [text]}), self.path('xlsx'))
        workbook = openpyxl.load_workbook(self.path('xlsx'), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual((workbook.active['A2'].value, workbook.active['A2'].data_type), (text, 's'))
        for frame in (pd.DataFrame({'Worker': [text + 'x']}), pd.DataFrame({text + 'x': ['worker']})):
            with self.assertRaises(ValueError):
                write_spreadsheet(frame, self.path('xlsx'))


if __name__ == '__main__':
    unittest.main()
